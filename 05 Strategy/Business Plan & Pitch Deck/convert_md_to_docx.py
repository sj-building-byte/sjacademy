import os
import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Convert Markdown to HTML
    html = markdown.markdown(md_text)
    soup = BeautifulSoup(html, 'html.parser')

    # Create a new Word document
    doc = Document()

    # Parse HTML and add to document
    for element in soup:
        if element.name == 'h1':
            doc.add_heading(element.text, level=1)
        elif element.name == 'h2':
            doc.add_heading(element.text, level=2)
        elif element.name == 'h3':
            doc.add_heading(element.text, level=3)
        elif element.name == 'p':
            p = doc.add_paragraph()
            # Handle basic bold/italic within paragraphs
            for child in element.children:
                if child.name == 'strong' or child.name == 'b':
                    p.add_run(child.text).bold = True
                elif child.name == 'em' or child.name == 'i':
                    p.add_run(child.text).italic = True
                else:
                    if child.string:
                        p.add_run(child.string)
        elif element.name == 'ul':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Bullet')
                for child in li.children:
                    if child.name == 'strong' or child.name == 'b':
                        p.add_run(child.text).bold = True
                    elif child.name == 'em' or child.name == 'i':
                        p.add_run(child.text).italic = True
                    else:
                        if child.string:
                            p.add_run(child.string)
        elif element.name == 'ol':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Number')
                for child in li.children:
                    if child.name == 'strong' or child.name == 'b':
                        p.add_run(child.text).bold = True
                    elif child.name == 'em' or child.name == 'i':
                        p.add_run(child.text).italic = True
                    else:
                        if child.string:
                            p.add_run(child.string)

    doc.save(docx_path)

# Convert all three files
files = [
    "SJ_Academy_Business_Plan.md",
    "SJ_Academy_Marketing_Plan.md",
    "SJ_Academy_Pitch_Deck.md"
]

for file in files:
    if os.path.exists(file):
        docx_name = file.replace('.md', '.docx')
        print(f"Converting {file} to {docx_name}...")
        md_to_docx(file, docx_name)
        print("Done.")
