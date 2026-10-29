from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("sj-academy-business-plan-draft.md")
OUTPUT = Path("sj-academy-business-plan-google-doc-ready.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def clean_inline(text):
    return text.replace("**", "").replace("`", "")


def split_table_row(line):
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [clean_inline(cell.strip()) for cell in stripped.split("|")]


def is_separator_row(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def add_markdown_table(doc, rows):
    parsed = [split_table_row(row) for row in rows if not is_separator_row(row)]
    if not parsed:
        return
    col_count = max(len(row) for row in parsed)
    parsed = [row + [""] * (col_count - len(row)) for row in parsed]
    table = doc.add_table(rows=len(parsed), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    set_table_width(table)
    set_cell_margins(table)
    set_table_borders(table)

    for r_idx, row in enumerate(parsed):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(9 if col_count >= 3 else 10)
            if r_idx == 0:
                run.bold = True
                set_cell_shading(cell, "F1F3F4")
            if c_idx > 0 and re.search(r"\d|PKR|%|M$", value):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after, color in [
        ("Heading 1", 20, 20, 6, RGBColor(0, 0, 0)),
        ("Heading 2", 16, 18, 6, RGBColor(0, 0, 0)),
        ("Heading 3", 14, 16, 4, RGBColor(67, 67, 67)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def add_paragraph_with_bold(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(clean_inline(part))
            run.bold = True
        else:
            run = paragraph.add_run(clean_inline(part))
        run.font.name = "Arial"
    return paragraph


def build_docx():
    text = SOURCE.read_text()
    doc = Document()
    style_document(doc)

    lines = text.splitlines()
    table_buffer = []
    in_code = False

    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            add_paragraph_with_bold(doc, line)
            continue

        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            table_buffer.append(line)
            continue

        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

        stripped = line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            title_para = doc.add_paragraph()
            title_para.paragraph_format.space_after = Pt(3)
            title_run = title_para.add_run(clean_inline(stripped[2:]))
            title_run.font.name = "Arial"
            title_run.font.size = Pt(26)
            title_run.bold = False
            continue

        if stripped.startswith("## "):
            doc.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
            continue

        if stripped.startswith("### "):
            doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            continue

        if stripped.startswith("- "):
            p = add_paragraph_with_bold(doc, stripped[2:], style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            continue

        numbered = re.match(r"^\d+\.\s+(.*)", stripped)
        if numbered:
            p = add_paragraph_with_bold(doc, numbered.group(1), style="List Number")
            p.paragraph_format.space_after = Pt(4)
            continue

        add_paragraph_with_bold(doc, stripped)

    if table_buffer:
        add_markdown_table(doc, table_buffer)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
